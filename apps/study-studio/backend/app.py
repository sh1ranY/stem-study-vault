import io
import json
import os
import hashlib
import sqlite3
import threading
import uuid
import zipfile
from datetime import datetime,timezone
from pathlib import Path
from typing import Optional, Dict, List

from dotenv import load_dotenv
from fastapi import FastAPI,HTTPException,UploadFile,File,Request
from fastapi.responses import Response,FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from pydantic import BaseModel,Field,ConfigDict
from .content import prepared_unit,SOURCES,checkpoint,transfer,public_task,TUTOR
from .assessment import assess
from . import provider

ROOT=Path(__file__).resolve().parents[1]
load_dotenv(ROOT/'.env')
DATA=Path(os.environ.get('STUDIO_DATA_DIR',str(ROOT/'data')))
DATA.mkdir(parents=True,exist_ok=True)
LOCK=threading.RLock()
MODEL_LOCK=threading.Lock()
MAX_UPLOAD=10*1024*1024

def now(): return datetime.now(timezone.utc).isoformat()
def uid(): return uuid.uuid4().hex
def connect():
    db=sqlite3.connect(str(DATA/'studio.sqlite3'),timeout=15)
    db.execute('PRAGMA journal_mode=WAL')
    db.execute('CREATE TABLE IF NOT EXISTS records (kind TEXT, id TEXT, body TEXT, PRIMARY KEY(kind,id))')
    return db
def put(kind,record):
    with connect() as db: db.execute('INSERT OR REPLACE INTO records VALUES (?,?,?)',(kind,record['id'],json.dumps(record,ensure_ascii=False)))
def get(kind,id):
    with connect() as db: row=db.execute('SELECT body FROM records WHERE kind=? AND id=?',(kind,id)).fetchone()
    if not row: raise HTTPException(404,'This record no longer exists. Return to your library.')
    return json.loads(row[0])
def all_records(kind):
    with connect() as db: rows=db.execute('SELECT body FROM records WHERE kind=?',(kind,)).fetchall()
    return [json.loads(row[0]) for row in rows]
def event(session,kind,**values):
    item={'id':uid(),'kind':kind,'at':now(),**values};session['events'].append(item);session['updated_at']=now();return item
def exposed(session,task_id):
    return any(e.get('task_id')==task_id and e['kind'] in ('hint','solution','attempt') for e in session['events'])
def public_session(session):
    data=json.loads(json.dumps(session))
    data['tasks']=[public_task(t) for t in session['tasks']]
    return data

app=FastAPI(title='STEM Study Vault · Learning Studio',docs_url='/api/docs',redoc_url=None)
app.add_middleware(TrustedHostMiddleware,allowed_hosts=['localhost','127.0.0.1','testserver'])

@app.middleware('http')
async def local_origin(request:Request,call_next):
    origin=request.headers.get('origin')
    if request.url.path.startswith('/api') and origin and origin not in ('http://127.0.0.1:8765','http://localhost:8765','http://127.0.0.1:5177','http://localhost:5177','http://testserver'):
        return Response('This local prototype only accepts its own browser origin.',status_code=403)
    if request.url.path.startswith('/api') and request.headers.get('sec-fetch-site')=='cross-site':
        return Response('Cross-site requests are not supported.',status_code=403)
    # Once pilot accounts exist, legacy prototype evidence is restricted to the owner.
    if request.url.path.startswith('/api/') and not request.url.path.startswith('/api/pilot/') and request.url.path not in ('/api/config','/api/docs','/api/openapi.json'):
        from .pilot import store as pilot_store
        from .pilot.routes import user as pilot_user
        if pilot_store.listing('user','auth'):
            try:
                if pilot_user(request)['role']!='owner': return Response('Owner access required for legacy prototype data.',status_code=403)
            except HTTPException:
                return Response('Sign in as the course owner to access legacy prototype data.',status_code=401)
    response=await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Referrer-Policy']='no-referrer'
    if request.url.path.startswith('/api'): response.headers['Cache-Control']='no-store'
    return response

class StrictModel(BaseModel):
    model_config=ConfigDict(extra='forbid')
class Create(StrictModel):
    mode:str='prepared'
    source_ids:List[str]=Field(default_factory=list,max_length=3)
class Progress(StrictModel):
    section_id:Optional[str]=None
    read_section:Optional[str]=None
    notes:Optional[str]=Field(default=None,max_length=30000)
class Attempt(StrictModel):
    task_id:str
    answers:Dict[str,str]
    choice:int
    working:str=Field(min_length=1,max_length=5000)
    confidence:str='unsure'
class Support(StrictModel):
    task_id:str
    level:str='hint'
class Ask(StrictModel):
    topic:str
class Correction(StrictModel):
    attempt_id:str
    reason:str=Field(min_length=1,max_length=2000)

@app.get('/api/config')
def config():
    return {'mode':'local-prototype','model_available':provider.available(),'provider':provider.configuration()[0] if provider.available() else None,
            'storage':'On this computer','max_upload_mb':10,'supported_files':['pdf','pptx','md','txt'],
            'notice':'Prepared ELEC270 uses authored teaching and rule-based checks. No learner study or live AI inference is implied.'}

@app.get('/api/sessions')
def sessions():
    return [{'id':s['id'],'title':s['unit']['title'],'course':s['unit']['course'],'updated_at':s['updated_at'],'read_count':len(s['read_sections']),'section_count':len(s['unit']['sections']),
             'attempt_count':len([a for a in s['attempts'] if a['active']]),'mode':s['mode']} for s in sorted(all_records('session'),key=lambda s:s['updated_at'],reverse=True)]

def extract(name,data):
    suffix=Path(name).suffix.lower();pages=[];warnings=[]
    if suffix in ('.md','.txt'):
        try: text=data.decode('utf-8-sig')
        except UnicodeDecodeError: raise HTTPException(422,'Please use UTF-8 text, a PDF or a PPTX file.')
        pages=[{'page':1,'text':text}]
    elif suffix=='.pdf':
        from pypdf import PdfReader
        try:
            reader=PdfReader(io.BytesIO(data))
            if reader.is_encrypted: raise ValueError('Encrypted PDF')
            if len(reader.pages)>60: raise HTTPException(422,'Please split this into a single unit of 60 pages or fewer.')
            pages=[{'page':i+1,'text':p.extract_text() or ''} for i,p in enumerate(reader.pages)]
        except HTTPException: raise
        except Exception: raise HTTPException(422,'This PDF could not be read. Export an unlocked PDF or upload text.')
        warnings.append('PDF text extraction may miss diagrams and mathematical layout. Compare the extracted text with the original before generation.')
    elif suffix=='.pptx':
        from pptx import Presentation
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist())>40*1024*1024: raise ValueError('Expanded file too large')
            prs=Presentation(io.BytesIO(data))
            if len(prs.slides)>60: raise HTTPException(422,'Please split this into a unit of 60 slides or fewer.')
            for i,slide in enumerate(prs.slides):
                parts=[]
                for shape in slide.shapes:
                    if shape.has_text_frame: parts.append(shape.text)
                    if shape.has_table: parts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                pages.append({'page':i+1,'text':'\n'.join(parts)})
        except HTTPException: raise
        except Exception: raise HTTPException(422,'This presentation could not be read. Try exporting one unit as PDF.')
        warnings.append('Text and tables were extracted. Images, speaker notes and embedded equations may need manual review.')
    else: raise HTTPException(415,'Supported formats: PDF, PPTX, Markdown and UTF-8 text.')
    if sum(len(p['text']) for p in pages)>60000: raise HTTPException(422,'The extracted text is too long for this prototype. Split it into a smaller unit.')
    empty=[p['page'] for p in pages if len(p['text'].strip())<20]
    if empty: warnings.append('Little or no text was found on pages '+', '.join(map(str,empty))+'. Scans and image-only equations are not OCR-processed.')
    return pages,warnings

@app.post('/api/imports')
async def upload(file:UploadFile=File(...)):
    data=await file.read(MAX_UPLOAD+1)
    if not data: raise HTTPException(422,'The selected file is empty.')
    if len(data)>MAX_UPLOAD: raise HTTPException(413,'The file is larger than 10 MB. Upload one smaller unit.')
    name=(file.filename or 'source').replace('\\','/').split('/')[-1][:180]
    pages,warnings=extract(name,data)
    record={'id':uid(),'name':name,'pages':pages,'warnings':warnings,'sha256':hashlib.sha256(data).hexdigest(),'created_at':now()}
    (DATA/(record['id']+Path(name).suffix.lower())).write_bytes(data)
    put('import',record)
    return record

@app.get('/api/imports/{id}/original')
def original(id:str):
    record=get('import',id)
    return FileResponse(DATA/(id+Path(record['name']).suffix.lower()),filename=record['name'],media_type='application/octet-stream')

@app.post('/api/sessions')
def create_session(body:Create):
    if body.mode not in ('prepared','model'): raise HTTPException(422,'Unknown teaching mode.')
    attached=[get('import',id) for id in body.source_ids]
    if body.mode=='model':
        if not attached: raise HTTPException(422,'Upload a source unit first.')
        sources=[{'id':f'{r["id"]}-p{p["page"]}','name':r['name'],'locator':f'Page {p["page"]}','text':p['text']} for r in attached for p in r['pages'] if p['text'].strip()]
        if sum(len(s['text'].strip()) for s in sources)<100: raise HTTPException(422,'Not enough readable text. Upload text or an accessible PDF.')
        if not provider.available(): raise HTTPException(409,'No model is configured. Your uploaded material is preserved. You can inspect it or open the separate prepared ELEC270 unit.')
        if not MODEL_LOCK.acquire(blocking=False): raise HTTPException(409,'Another generation is running. Please wait before starting another.')
        run={'id':uid(),'at':now(),'operation':'generate','status':'started','call_limit':1};put('run',run)
        try:
            unit,usage=provider.generate(sources)
            run.update(status='completed',usage=usage);put('run',run)
        except Exception as exc:
            run.update(status='failed');put('run',run)
            message=str(exc) if isinstance(exc,ValueError) else 'The provider request failed or timed out. No automatic retry was made.'
            raise HTTPException(502,message)
        finally: MODEL_LOCK.release()
        tasks=[];mode='model-draft'
    else:
        unit=prepared_unit();sources=SOURCES;tasks=[checkpoint()];mode='prepared'
    session={'id':uid(),'mode':mode,'unit':unit,'sources':sources,'attached_sources':attached,'tasks':tasks,'attempts':[],'events':[],
             'read_sections':[],'current_section':unit['sections'][0]['id'],'notes':'','created_at':now(),'updated_at':now(),'transfer_count':0,'tutor_messages':[]}
    event(session,'created',mode=mode)
    put('session',session)
    return public_session(session)

@app.get('/api/sessions/{id}')
def read_session(id:str): return public_session(get('session',id))

@app.patch('/api/sessions/{id}')
def progress(id:str,body:Progress):
    with LOCK:
        s=get('session',id);valid={x['id'] for x in s['unit']['sections']}
        for field in (body.section_id,body.read_section):
            if field is not None and field not in valid: raise HTTPException(422,'Unknown lesson section.')
        if body.section_id: s['current_section']=body.section_id
        if body.read_section and body.read_section not in s['read_sections']:
            s['read_sections'].append(body.read_section);event(s,'read',section_id=body.read_section)
        if body.notes is not None: s['notes']=body.notes
        s['updated_at']=now();put('session',s);return public_session(s)

def find_task(session,id):
    for task in session['tasks']:
        if task['id']==id: return task
    raise HTTPException(404,'This task is not part of the current unit.')

@app.post('/api/sessions/{id}/attempts')
def attempt(id:str,body:Attempt):
    if not body.working.strip(): raise HTTPException(422,'Add one key step. A short calculation is enough.')
    if body.confidence not in ('unsure','partly','confident'): raise HTTPException(422,'Choose a confidence level.')
    with LOCK:
        s=get('session',id);task=find_task(s,body.task_id)
        seen=exposed(s,task['id']);prior=len([a for a in s['attempts'] if a['task_id']==task['id']])
        try: result=assess(task,body.answers,body.choice,seen,prior)
        except ValueError as exc: raise HTTPException(422,str(exc))
        ev=event(s,'attempt',task_id=task['id'],exposure_before=seen)
        record={'id':ev['id'],'task_id':task['id'],'at':ev['at'],'answers':body.answers,'choice':body.choice,'working':body.working,'confidence':body.confidence,
                'exposure_before':seen,'result':result,'active':True,'origin':'user-submitted'}
        s['attempts'].append(record);put('session',s);return public_session(s)

@app.post('/api/sessions/{id}/support')
def support(id:str,body:Support):
    if body.level not in ('hint','solution'): raise HTTPException(422,'Choose a hint or solution.')
    with LOCK:
        s=get('session',id);task=find_task(s,body.task_id)
        ev=event(s,body.level,task_id=task['id']);put('session',s)
        return {'body':task[body.level],'event_id':ev['id'],'session':public_session(s)}

@app.post('/api/sessions/{id}/fresh-task')
def fresh(id:str):
    with LOCK:
        s=get('session',id)
        if s['mode']!='prepared': raise HTTPException(409,'Generated drafts do not yet have a reviewed assessment pack. Use prepared ELEC270 to demonstrate the full learning loop.')
        task=transfer(s['transfer_count'])
        if task is None: raise HTTPException(409,'All six prepared variants have been opened. They will not be recycled as unseen tasks. Export your record or ask a teacher for a new task.')
        s['transfer_count']+=1;s['tasks'].append(task);event(s,'task_opened',task_id=task['id']);put('session',s)
        return public_session(s)

@app.post('/api/sessions/{id}/tutor')
def tutor(id:str,body:Ask):
    if body.topic not in TUTOR: raise HTTPException(422,'Choose one of the prepared help topics.')
    with LOCK:
        s=get('session',id)
        if s['mode']!='prepared': raise HTTPException(409,'Prepared signal support is separate from this uploaded lesson.')
        entry={'id':uid(),'at':now(),**TUTOR[body.topic],'mode':'prepared'};s['tutor_messages'].append(entry)
        # General tutor help while a task is open conservatively counts as support.
        for task in s['tasks']:
            event(s,'hint',task_id=task['id'],via='tutor',topic=body.topic)
        put('session',s)
        return public_session(s)

@app.post('/api/sessions/{id}/corrections')
def correct(id:str,body:Correction):
    if not body.reason.strip(): raise HTTPException(422,'Give a short reason for the correction.')
    with LOCK:
        s=get('session',id)
        target=next((a for a in s['attempts'] if a['id']==body.attempt_id),None)
        if target is None: raise HTTPException(404,'Attempt not found.')
        target['active']=False;target['withdrawal_reason']=body.reason
        event(s,'correction',attempt_id=body.attempt_id,reason=body.reason)
        put('session',s);return public_session(s)

@app.delete('/api/sessions/{id}')
def delete(id:str):
    with LOCK:
        get('session',id)
        with connect() as db: db.execute('DELETE FROM records WHERE kind=? AND id=?',('session',id))
    return {'deleted':True}

@app.get('/api/sessions/{id}/export')
def export(id:str):
    with LOCK:
        s=get('session',id)
        for task in s['tasks']:
            event(s,'solution',task_id=task['id'],via='export')
        put('session',s)
    lesson='# '+s['unit']['title']+'\n\n'+s['unit']['summary']+'\n\n'+s['unit']['scope']+'\n\n'
    lesson+='\n\n'.join('## '+x['title']+'\n\n'+x['body']+('\n\n![Ramp comparison](Assets/ramp-comparison.svg)\n' if x.get('figure')=='ramps' else '')+'\n\nSources: '+', '.join(x['source_refs']) for x in s['unit']['sections'])
    practice='# Practice and support\n\n'
    for t in s['tasks']:
        practice+='## '+t['title']+'\n\n'+t['prompt']+'\n\n'+t.get('template','')+'\n\n'
        practice+='\n'.join('- '+f['label'] for f in t['fields'])+'\n\n'
        practice+='\n'.join(f'{i+1}. '+c for i,c in enumerate(t['choices']))+'\n\n'
        for label in ('hint','solution'):
            practice+=f'> [!{label if label=="hint" else "success"}]- {label.title()}\n'+ '\n'.join('> '+line for line in t[label].splitlines())+'\n\n'
    record='# Learning evidence\n\nPrepared-mode checks cover numerical fields and reasoning choices. Free working is stored, not semantically graded. Delayed retention is unassessed.\n\n'
    for a in s['attempts']:
        record+=f'## {a["at"]} · {a["result"]["status"]}'+(' (withdrawn)' if not a['active'] else '')+'\n\n'+a['working']+'\n\n'+a['result']['support']['hypothesis']+'\n\n'
    sources='# Sources\n\n'+ '\n\n'.join(f'## {r["id"]} · {r["name"]}\n\n{r["locator"]}\n\n{r["text"]}' for r in s['sources'])
    blob=io.BytesIO()
    with zipfile.ZipFile(blob,'w',zipfile.ZIP_DEFLATED) as z:
        for name,body in {'Start.md':'# Learning Studio export\n\n[[Lesson]] · [[Practice]] · [[Evidence]] · [[Sources]] · [[Personal notes]]\n\nOriginal teaching sample; not an official course handout.\n',
                          'Lesson.md':lesson,'Practice.md':practice,'Evidence.md':record,'Sources.md':sources,'Personal notes.md':s['notes'] or '# Personal notes\n'}.items(): z.writestr(name,body)
        if s['mode']=='prepared': z.writestr('Assets/ramp-comparison.svg',(ROOT/'backend'/'ramp-comparison.svg').read_text())
        z.writestr('Maintenance/learning-record.json',json.dumps(s,ensure_ascii=False,indent=2))
        for r in s['attached_sources']:
            z.writestr('Sources/Extracted/'+r['id']+'.json',json.dumps(r,ensure_ascii=False,indent=2))
            raw=DATA/(r['id']+Path(r['name']).suffix.lower())
            if raw.exists(): z.write(raw,'Sources/Original/'+r['id'][:8]+'-'+r['name'])
    return Response(blob.getvalue(),media_type='application/zip',headers={'Content-Disposition':'attachment; filename="STEM-Study-Vault-'+id[:8]+'.zip"'})

@app.get('/api/sample-source')
def sample_source():
    body='# Original signal construction brief\n\nELEC270-aligned authored demonstration; not official lecture material.\n\n'
    body+='\n\n'.join('## '+s['locator']+'\n\n'+s['text'] for s in SOURCES)
    return Response(body,media_type='text/markdown',headers={'Content-Disposition':'attachment; filename="signal-construction-brief.md"'})

from .pilot.routes import router as pilot_router
app.include_router(pilot_router)

if (ROOT/'dist').exists():
    app.mount('/',StaticFiles(directory=ROOT/'dist',html=True),name='frontend')
