"""Local integration tests with isolated data; never makes a paid model call."""
import copy
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
import zipfile
from unittest.mock import patch,MagicMock

TEMP=tempfile.TemporaryDirectory()
os.environ['STUDIO_DATA_DIR']=TEMP.name
from fastapi.testclient import TestClient
from backend import app as studio,provider
from backend.assessment import numeric
from backend.content import transfer

class StudioTests(unittest.TestCase):
    def setUp(self):
        from backend.pilot import store
        from sqlalchemy import delete
        with store.engine.begin() as connection: connection.execute(delete(store.records))
        self.client=TestClient(studio.app)
        self.env=patch.dict(os.environ,{'STUDIO_PROVIDER':'','STUDIO_MODEL':'','OPENAI_API_KEY':'','DEEPSEEK_API_KEY':''});self.env.start()
        with studio.connect() as db: db.execute('DELETE FROM records')
        self.s=self.client.post('/api/sessions',json={'mode':'prepared'}).json()
        self.url='/api/sessions/'+self.s['id']
    def tearDown(self): self.env.stop();self.client.close()
    def attempt(self,answers=None,choice=1,working='5 - 2 = 3, then 5 - 3 = 2',task=None):
        return self.client.post(self.url+'/attempts',json={'task_id':task or self.s['tasks'][0]['id'],'answers':answers or {'shift':'3','tail3':'2','tail5':'2'},'choice':choice,'working':working,'confidence':'partly'})
    def fresh(self):
        r=self.client.post(self.url+'/fresh-task');self.assertEqual(r.status_code,200);return r.json()['tasks'][-1]['id']
    def correct_transfer(self,task=None):
        return self.attempt({'a':'2','b':'-2','c':'-6','end':'0'},2,'The plateau is 2*(4-1)=6, so remove six.',task or self.fresh())
    def test_answers_redacted_and_checkpoint_supported(self):
        text=json.dumps(self.s['tasks']);self.assertNotIn('answer',text);self.assertNotIn('correct_choice',text)
        result=self.attempt().json()['attempts'][-1]['result']
        self.assertTrue(result['correct']);self.assertEqual(result['status'],'supported')
        self.assertEqual(result['retention'],'unassessed')
    def test_different_errors_choose_different_teaching(self):
        r=self.attempt({'shift':'5','tail3':'2','tail5':'2'}).json()['attempts'][-1]['result'];self.assertEqual(r['strategy'],'substitution')
        r=self.attempt({'shift':'3','tail3':'0','tail5':'0'},0).json()['attempts'][-1]['result'];self.assertEqual(r['strategy'],'height')
        r=self.attempt(choice=0).json()['attempts'][-1]['result'];self.assertEqual(r['strategy'],'reasoning')
    def test_independent_new_task_then_supported_retry(self):
        task=self.fresh()
        r=self.correct_transfer(task);self.assertEqual(r.status_code,200)
        self.assertEqual(r.json()['attempts'][-1]['result']['status'],'independent')
        self.assertEqual(self.correct_transfer(task).json()['attempts'][-1]['result']['status'],'supported')
    def test_hint_solution_and_tutor_exposure(self):
        for route,payload in [('/support',{'level':'hint'}),('/support',{'level':'solution'}),('/tutor',{'topic':'height'})]:
            with self.subTest(route=route,payload=payload):
                self.setUpSession()
                task=self.fresh()
                if route=='/support':payload={**payload,'task_id':task}
                self.assertEqual(self.client.post(self.url+route,json=payload).status_code,200)
                self.assertEqual(self.correct_transfer(task).json()['attempts'][-1]['result']['status'],'supported')
    def setUpSession(self):
        self.s=self.client.post('/api/sessions',json={}).json();self.url='/api/sessions/'+self.s['id']
    def test_general_support_covers_older_open_tasks(self):
        task=self.fresh();self.fresh();self.client.post(self.url+'/tutor',json={'topic':'construction'})
        self.assertEqual(self.correct_transfer(task).json()['attempts'][-1]['result']['status'],'supported')
    def test_withdraw_does_not_erase_exposure(self):
        task=self.fresh();a=self.correct_transfer(task).json()['attempts'][-1]
        r=self.client.post(self.url+'/corrections',json={'attempt_id':a['id'],'reason':'Entered a copied answer by mistake.'}).json()
        self.assertFalse(r['attempts'][-1]['active']);self.assertTrue(any(e['kind']=='correction' for e in r['events']))
        self.assertEqual(self.correct_transfer(task).json()['attempts'][-1]['result']['status'],'supported')
        self.assertEqual(self.client.post(self.url+'/corrections',json={'attempt_id':a['id'],'reason':'   '}).status_code,422)
    def test_six_distinct_variants_not_recycled(self):
        tasks=[self.fresh() for _ in range(6)]
        self.assertEqual(len(set(tasks)),6)
        self.assertEqual(self.client.post(self.url+'/fresh-task').status_code,409)
        internal=studio.get('session',self.s['id'])['tasks'][1:]
        self.assertEqual(len(set(t['prompt'] for t in internal)),6)
    def test_incomplete_invalid_answers_cannot_be_observations(self):
        for answers,choice,working in [({'shift':'3'},1,'step'),({'shift':'__import__("os")','tail3':'2','tail5':'2'},1,'step'),({'shift':'3','tail3':'2','tail5':'2'},-1,'step'),({'shift':'3','tail3':'2','tail5':'2'},1,' ')]:
            self.assertEqual(self.attempt(answers,choice,working).status_code,422)
        self.assertEqual(len(self.client.get(self.url).json()['attempts']),0)
    def test_arithmetic_parser(self):
        self.assertEqual(numeric('−6/2'),-3);self.assertEqual(numeric('2 × (4 - 1)'),6)
        for text in ['1/0','1e999','2**10000000','open("x")','True','', '[1][0]']:
            with self.assertRaises(ValueError): numeric(text)
    def test_persist_read_notes_and_export_with_original(self):
        imported=self.client.post('/api/imports',files={'file':('brief.md',b'# Original\nA ramp is zero before time zero.','text/markdown')}).json()
        self.s=self.client.post('/api/sessions',json={'source_ids':[imported['id']]}).json();self.url='/api/sessions/'+self.s['id']
        self.client.patch(self.url,json={'read_section':'shifts','section_id':'plateau','notes':'My own explanation.'})
        with TestClient(studio.app) as other:
            s=other.get(self.url).json()
        self.assertEqual(s['current_section'],'plateau');self.assertEqual(s['notes'],'My own explanation.');self.assertEqual(s['read_sections'],['shifts'])
        self.attempt()
        r=self.client.get(self.url+'/export');self.assertEqual(r.status_code,200)
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            self.assertEqual(z.read('Personal notes.md').decode(),'My own explanation.')
            self.assertIn('Assets/ramp-comparison.svg',z.namelist())
            self.assertIn('Assets/ramp-comparison.svg',z.read('Lesson.md').decode())
            self.assertTrue(any(p.startswith('Sources/Original/') for p in z.namelist()))
            self.assertEqual(len(json.loads(z.read('Maintenance/learning-record.json'))['attempts']),1)
    def test_upload_rejections_and_empty_pdf_notice(self):
        for name,data,status in [('bad.exe',b'content',415),('empty.txt',b'',422),('invalid.txt',b'\xff',422),('bad.pdf',b'invalid',422),('large.txt',b'a'*(studio.MAX_UPLOAD+1),413),('too-long.txt',b'a'*60001,422)]:
            self.assertEqual(self.client.post('/api/imports',files={'file':(name,data)}).status_code,status)
        from pypdf import PdfWriter
        b=io.BytesIO();w=PdfWriter();w.add_blank_page(width=100,height=100);w.write(b)
        r=self.client.post('/api/imports',files={'file':('scan.pdf',b.getvalue())});self.assertEqual(r.status_code,200)
        self.assertTrue(any('OCR' in w for w in r.json()['warnings']))
        self.assertEqual(self.client.post('/api/sessions',json={'mode':'model','source_ids':[r.json()['id']]}).status_code,422)
    def test_pptx_text_extraction(self):
        from pptx import Presentation
        p=Presentation();slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text='Original signal unit'
        b=io.BytesIO();p.save(b)
        r=self.client.post('/api/imports',files={'file':('unit.pptx',b.getvalue())});self.assertEqual(r.status_code,200)
        self.assertIn('Original signal unit',r.json()['pages'][0]['text'])
    def test_no_key_keeps_upload_without_model_call(self):
        r=self.client.post('/api/imports',files={'file':('unit.md',b'Original source. '*20)}).json()
        with patch.object(provider,'generate') as generate:
            self.assertEqual(self.client.post('/api/sessions',json={'mode':'model','source_ids':[r['id']]}).status_code,409);generate.assert_not_called()
        self.assertEqual(self.client.get('/api/imports/'+r['id']+'/original').status_code,200)
        self.assertFalse(self.client.get('/api/config').json()['model_available'])
    def test_cross_origin_rejected(self):
        self.assertEqual(self.client.post('/api/sessions',json={},headers={'Origin':'https://example.com'}).status_code,403)
    def test_provider_schema_rejects_unsafe_structures(self):
        value={'title':'A unit','summary':'Summary','prerequisites':'Algebra','objectives':['Evaluate a ramp'],'sections':[{'title':'Shift','body':'A ramp is shifted by substituting the complete argument. '*3,'source_refs':['source-1'],'figure':'ramps'}]}
        clean=provider.validate_unit(copy.deepcopy(value),{'source-1'});self.assertNotIn('figure',clean['sections'][0])
        for mutation in [{'prerequisites':{}},{'objectives':None},{'sections':[{'title':'x','body':'x'*120,'source_refs':[{}]}]}]:
            with self.assertRaises(ValueError):provider.validate_unit({**value,**mutation},{'source-1'})
    def test_mocked_model_contract_and_failure_no_retry(self):
        value={'title':'Signals','summary':'A generated draft','prerequisites':'Basic algebra','objectives':['Substitute correctly'],'sections':[{'title':'Shift','body':'A ramp is evaluated by substituting the entire argument. '*3,'source_refs':['source-1']}]}
        with patch.dict(os.environ,{'STUDIO_PROVIDER':'openai','OPENAI_API_KEY':'test-not-a-real-key','STUDIO_MODEL':'configured-model'}):
            with patch.object(provider.httpx,'Client') as client:
                response=client.return_value.__enter__.return_value.post.return_value
                response.status_code=200;response.json.return_value={'choices':[{'finish_reason':'stop','message':{'content':json.dumps(value)}}],'usage':{'total_tokens':100}}
                unit,usage=provider.generate([{'id':'source-1','text':'source'}]);self.assertEqual(unit['title'],'Signals');self.assertEqual(usage['total_tokens'],100)
                call=client.return_value.__enter__.return_value.post;self.assertEqual(call.call_count,1)
                self.assertNotIn('test-not-a-real-key',self.client.get('/api/config').text)
                response.status_code=429
                with self.assertRaises(ValueError):provider.generate([{'id':'source-1','text':'source'}])
                self.assertEqual(call.call_count,2)
    def test_export_answers_count_as_exposure(self):
        task=self.fresh()
        self.assertEqual(self.client.get(self.url+'/export').status_code,200)
        self.assertEqual(self.correct_transfer(task).json()['attempts'][-1]['result']['status'],'supported')
    def test_generated_draft_has_no_invented_assessments(self):
        r=self.client.post('/api/imports',files={'file':('unit.md',b'Original source. '*20)}).json()
        lesson={'title':'Draft','summary':'Draft summary','prerequisites':'Algebra','objectives':['Substitute'],'sections':[{'title':'An idea','body':'A long explanation. '*10,'source_refs':[r['id']+'-p1']}]}
        validated=provider.validate_unit(lesson,{r['id']+'-p1'})
        with patch.object(provider,'available',return_value=True),patch.object(provider,'generate',return_value=(validated,{'total_tokens':100})) as generation:
            response=self.client.post('/api/sessions',json={'mode':'model','source_ids':[r['id']]});self.assertEqual(response.status_code,200)
            draft=response.json();self.assertEqual(draft['mode'],'model-draft');self.assertEqual(draft['tasks'],[]);generation.assert_called_once()
            self.assertEqual(self.client.post('/api/sessions/'+draft['id']+'/fresh-task').status_code,409)
    def test_bad_read_does_not_write(self):
        self.assertEqual(self.client.patch(self.url,json={'read_section':'unknown','notes':'should not save'}).status_code,422)
        self.assertEqual(self.client.get(self.url).json()['notes'],'')

if __name__=='__main__':unittest.main()
