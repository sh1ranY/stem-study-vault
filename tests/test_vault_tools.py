import importlib.util
import json
import math
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/stem-study-vault/scripts/vault_tools.py"
spec = importlib.util.spec_from_file_location("vault_tools", SCRIPT)
tools = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tools)


class FileFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.vault = self.root / "vault"
        self.vault.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def write(self, relative, body=""):
        path = self.vault / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(body, encoding="utf-8")
        return path

    def run_cli(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)],
                              capture_output=True, text=True)


class VaultTests(FileFixture):
    def test_inventory_relative_hashes_and_hidden_exclusion(self):
        self.write("notes/中文 note.md", "hello")
        self.write(".secret", "private")
        first = tools.inventory(self.vault)
        self.assertEqual(first, tools.inventory(self.vault))
        self.assertEqual([r["path"] for r in first["files"]], ["notes/中文 note.md"])
        self.assertEqual(first["files"][0]["sha256"], tools.hashlib.sha256(b"hello").hexdigest())
        self.assertIn(".secret", first["skipped"])

    def test_symlink_is_not_followed(self):
        secret = self.root / "outside.txt"
        secret.write_text("not course material")
        try:
            (self.vault / "alias.txt").symlink_to(secret)
        except OSError:
            self.skipTest("Symlink creation not permitted")
        result = tools.inventory(self.vault)
        self.assertEqual(result["files"], [])
        self.assertEqual(result["skipped"], ["alias.txt"])

    def test_missing_root_has_nonzero_exit(self):
        result = self.run_cli("inventory", "--source", self.root / "missing")
        self.assertEqual(result.returncode, 2)

    def test_inventory_refuses_output_in_source(self):
        result = self.run_cli("inventory", "--source", self.vault,
                              "--output", self.vault / "manifest.json")
        self.assertEqual(result.returncode, 2)
        self.assertFalse((self.vault / "manifest.json").exists())

    def test_output_is_never_overwritten(self):
        output = self.root / "before.json"
        output.write_text("keep me")
        result = self.run_cli("snapshot", "--source", self.vault, "--output", output)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(output.read_text(), "keep me")

    def test_initialization_has_resolvable_navigation(self):
        tools.initialize(self.vault, "circuits", "Circuit course")
        self.assertTrue(tools.check_links(self.vault)["ok"])
        self.assertIn("scaffold", (self.vault / "circuits/Course.md").read_text())

    def test_initialization_preserves_existing_course(self):
        self.write("circuits/Personal/me.md", "My own thoughts")
        with self.assertRaises(ValueError):
            tools.initialize(self.vault, "circuits", "New")
        self.assertEqual((self.vault / "circuits/Personal/me.md").read_text(), "My own thoughts")

    def test_initialization_rejects_path_traversal(self):
        for course in ("../outside", "/absolute", "a/b", "..", "a\\b"):
            with self.assertRaises(ValueError):
                tools.initialize(self.vault, course, "Title")

    def test_snapshot_detects_changes_and_deletions(self):
        a = self.write("a.md", "one")
        b = self.write("b.md", "two")
        output = self.root / "before.json"
        tools.write_new_json(output, tools.inventory(self.vault))
        a.write_text("changed")
        b.unlink()
        self.write("c.md", "added")
        result = tools.compare_snapshot(self.vault, output)
        self.assertFalse(result["ok"])
        self.assertEqual(result["changed"], ["a.md"])
        self.assertEqual(result["missing"], ["b.md"])
        self.assertEqual(result["added"], ["c.md"])

    def test_snapshot_additions_do_not_mutate_existing_files(self):
        self.write("a.md", "one")
        output = self.root / "before.json"
        tools.write_new_json(output, tools.inventory(self.vault))
        self.write("b.md", "two")
        self.assertTrue(tools.compare_snapshot(self.vault, output)["ok"])

    def test_wiki_links_headings_blocks_assets_and_table_alias(self):
        self.write("A.md", "[[folder/B#电压与电流\\|B]]\n[[folder/B#^q-1]]\n![[asset.svg|400]]")
        self.write("folder/B.md", "# 电压与电流\nAnswer\n^q-1")
        self.write("asset.svg", "<svg/>")
        self.assertTrue(tools.check_links(self.vault)["ok"])

    def test_link_failures_are_not_silenced(self):
        self.write("A.md", "[[B#missing]]\n[[B#^no-id]]\n![[missing.png]]")
        self.write("B.md", "# Present")
        result = tools.check_links(self.vault)
        self.assertEqual(len(result["errors"]), 3)
        self.assertFalse(result["ok"])

    def test_ambiguous_basename_detected(self):
        self.write("A.md", "[[Shared]]")
        self.write("one/Shared.md")
        self.write("two/Shared.md")
        self.assertEqual(tools.check_links(self.vault)["errors"][0]["reason"], "Ambiguous target")

    def test_explicit_path_resolves_ambiguity(self):
        self.write("A.md", "[[two/Shared]]")
        self.write("one/Shared.md")
        self.write("two/Shared.md")
        self.assertTrue(tools.check_links(self.vault)["ok"])

    def test_local_heading_and_percent_encoded_target(self):
        self.write("A.md", "# **Start**\n[[#Start]]\n[[Space%20name]]")
        self.write("Space name.md")
        self.assertTrue(tools.check_links(self.vault)["ok"])

    def test_code_examples_are_not_live_links(self):
        self.write("A.md", "```markdown\n[[not-a-link]]\n```\n`[[also-not-a-link]]`\n")
        result = tools.check_links(self.vault)
        self.assertTrue(result["ok"])
        self.assertEqual(result["wiki_links_checked"], 0)

    def test_outside_wiki_target_rejected(self):
        self.write("A.md", "[[../outside.md]]")
        (self.root / "outside.md").write_text("outside")
        self.assertFalse(tools.check_links(self.vault)["ok"])

    def test_markdown_links_warn_about_check_limit(self):
        self.write("A.md", "[not checked](missing.md)")
        self.assertTrue(tools.check_links(self.vault)["warnings"])

    def test_extract_preserves_source_and_reports_unsupported(self):
        source = self.write("lecture.md", "# Heading\nEquation: x=2")
        self.write("figure.png", "image fixture")
        before = tools.sha256(source)
        result = tools.extract_tree(self.vault, self.root / "extracted")
        self.assertFalse(result["ok"])
        extracted = json.loads((self.root / "extracted/lecture.md.json").read_text())
        self.assertEqual(extracted["units"][0]["text"], source.read_text())
        self.assertEqual(tools.sha256(source), before)

    def test_extract_rejects_overlapping_and_existing_output(self):
        for output in (self.vault, self.vault / "out", self.root):
            with self.assertRaises(ValueError):
                tools.extract_tree(self.vault, output)

    def test_extract_empty_directory_is_not_successful_coverage(self):
        self.assertFalse(tools.extract_tree(self.vault, self.root / "out")["ok"])

    def test_cli_end_to_end(self):
        self.write("lecture.md", "# A\nSome teaching material")
        init = self.run_cli("init", "--vault", self.root / "student", "--course", "demo", "--title", "Demo")
        self.assertEqual(init.returncode, 0, init.stderr)
        extract = self.run_cli("extract", "--source", self.vault, "--output", self.root / "text")
        self.assertEqual(extract.returncode, 0, extract.stderr)
        check = self.run_cli("check", "--vault", self.root / "student")
        self.assertEqual(check.returncode, 0, check.stderr)
        repeat = self.run_cli("init", "--vault", self.root / "student", "--course", "demo", "--title", "Again")
        self.assertEqual(repeat.returncode, 2)


class OptionalExtractionTests(FileFixture):
    def test_pdf_page_boundaries_empty_pages_and_page_links(self):
        try:
            from pypdf import PdfWriter
            from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        except ImportError:
            self.skipTest("Install extraction dependencies to run PDF integration test")
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 150 Td (Voltage = 5 V) Tj ET")
        page[NameObject("/Contents")] = stream
        writer.add_blank_page(width=300, height=200)
        path = self.vault / "lecture.pdf"
        with path.open("wb") as file:
            writer.write(file)
        result = tools.extract_file(path)
        self.assertEqual([u["locator"] for u in result["units"]], ["physical-page-1", "physical-page-2"])
        self.assertIn("Voltage", result["units"][0]["text"])
        self.assertTrue(result["units"][1]["empty_text"])
        self.write("A.md", "[[lecture.pdf#page=2]]\n[[lecture.pdf#page=3]]")
        checked = tools.check_links(self.vault)
        self.assertEqual(len(checked["errors"]), 1)
        self.assertEqual(checked["errors"][0]["reason"], "PDF page out of range")

    def test_pptx_slide_boundaries_and_tables(self):
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            self.skipTest("Install extraction dependencies to run PPTX integration test")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Current I"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = "R"
        table.cell(0, 1).text = "1000 ohm"
        presentation.slides.add_slide(presentation.slide_layouts[6])
        path = self.vault / "slides.pptx"
        presentation.save(path)
        result = tools.extract_file(path)
        self.assertEqual(len(result["units"]), 2)
        self.assertIn("1000 ohm", result["units"][0]["text"])
        self.assertTrue(result["units"][1]["empty_text"])


class ExampleTests(unittest.TestCase):
    def test_demo_wiki_links(self):
        result = tools.check_links(ROOT / "examples/demo-vault")
        self.assertTrue(result["ok"], result["errors"])
        self.assertGreater(result["wiki_links_checked"], 10)

    def test_rc_worked_numbers_and_differential_equation(self):
        # Independent evaluation of the demonstrated physical model and exercises.
        r, c, vs = 1000.0, 100e-6, 5.0
        tau = r * c
        self.assertAlmostEqual(tau, 0.1)
        self.assertAlmostEqual(vs * (1 - math.exp(-1)), 3.160602794, places=8)
        self.assertAlmostEqual(-tau * math.log(0.1), 0.2302585093, places=9)
        self.assertAlmostEqual(12 * (1 - math.exp(-1)), 7.585446706, places=8)
        for t in (0.0, 0.02, 0.1, 0.5):
            voltage = vs * (1 - math.exp(-t / tau))
            derivative = vs / tau * math.exp(-t / tau)
            self.assertAlmostEqual(r * c * derivative + voltage, vs)
        self.assertAlmostEqual(5 + (2 - 5) * math.exp(-1), 3.896361676, places=8)


if __name__ == "__main__":
    unittest.main()
