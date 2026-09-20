#!/usr/bin/env python3
"""Local, non-overwriting helpers. Python 3.9+; extraction deps are optional."""

import argparse
import hashlib
import html
import json
import re
import sys
import unicodedata
from pathlib import Path
from urllib.parse import unquote


def root_dir(value):
    root = Path(value).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("Not a directory: {}".format(root))
    return root


def inside(path, root):
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def collect(root):
    """Skip hidden entries and symlinks, with an inspectable skipped list."""
    files, skipped = [], []

    def walk(folder):
        for path in sorted(folder.iterdir(), key=lambda p: p.name):
            rel = path.relative_to(root).as_posix()
            if path.is_symlink() or path.name.startswith(".") or path.name == "__pycache__":
                skipped.append(rel)
            elif path.is_dir():
                walk(path)
            elif path.is_file():
                files.append(path)
    walk(root)
    return files, skipped


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def inventory(root):
    files, skipped = collect(root)
    return {
        "schema_version": 1,
        "files": [{"path": p.relative_to(root).as_posix(), "bytes": p.stat().st_size,
                   "sha256": sha256(p)} for p in files],
        "skipped": skipped,
    }


def write_new_json(path, data):
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def outside_output(path, root):
    path = Path(path).expanduser()
    if inside(path, root):
        raise ValueError("Output must be outside the scanned source/protected tree")
    if path.exists() or path.is_symlink():
        raise ValueError("Output already exists; choose a new path")
    return path


def initialize(vault, course, title):
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}", course):
        raise ValueError("Course folder must be a short slug: letters, digits, _ or -")
    if not title.strip() or any(c in title for c in "\r\n[]"):
        raise ValueError("Use a non-empty, single-line title without square brackets")
    vault = Path(vault).expanduser().resolve()
    target = vault / course
    if target.exists() or target.is_symlink():
        raise ValueError("Course already exists; initialization never overwrites it")
    target.mkdir(parents=True)
    for folder in ("Lessons", "Topics", "Practice", "Personal", "Assets",
                   "Sources/Original", "Sources/Extracted", "Maintenance"):
        (target / folder).mkdir(parents=True, exist_ok=True)
    (target / "Course.md").write_text(
        "# {}\n\nThis course is a scaffold; teaching content has not been built.\n\n"
        "Start by providing sources and describing your learning goal.\n\n"
        "- [[{}/Learning profile|Learning profile]]\n"
        "- [[{}/Maintenance/Progress|Build progress]]\n".format(title, course, course),
        encoding="utf-8")
    (target / "Learning profile.md").write_text(
        "# Learning profile\n\nGoal: systematic learning (provisional).\n\n"
        "Background and language: not yet assessed.\n\n"
        "Personal notes: preserve Personal/ and any existing user-authored material.\n",
        encoding="utf-8")
    (target / "Maintenance/Progress.md").write_text(
        "# Build progress\n\nStatus: scaffold only.\n\nNext: inspect sources, "
        "establish learner preferences and build one complete unit.\n", encoding="utf-8")
    return {"created": str(target), "status": "scaffold-only"}


def extract_file(path):
    suffix = path.suffix.lower()
    warnings, units = [], []
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF extraction requires pypdf; see requirements-extract.txt") from exc
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("Encrypted PDF needs an accessible copy")
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            units.append({"locator": "physical-page-{}".format(number), "text": text,
                          "empty_text": not bool(text.strip())})
        warnings.append("Text only: inspect original formulas, figures and tables; no OCR performed.")
    elif suffix == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ValueError("PPTX extraction requires python-pptx; see requirements-extract.txt") from exc

        def shape_text(shape):
            if getattr(shape, "has_text_frame", False):
                return shape.text_frame.text
            if getattr(shape, "has_table", False):
                return "\n".join("\t".join(c.text for c in row.cells) for row in shape.table.rows)
            if hasattr(shape, "shapes"):
                return "\n".join(shape_text(s) for s in shape.shapes)
            return ""
        presentation = Presentation(str(path))
        for number, slide in enumerate(presentation.slides, 1):
            text = "\n".join(shape_text(shape) for shape in slide.shapes)
            units.append({"locator": "slide-{}".format(number), "text": text,
                          "empty_text": not bool(text.strip())})
        warnings.append("Includes slide positions, even hidden slides; excludes notes and image/math rendering.")
    elif suffix in (".md", ".txt"):
        text = path.read_text(encoding="utf-8-sig")
        units.append({"locator": "document", "text": text, "empty_text": not bool(text.strip())})
    else:
        raise ValueError("Unsupported format: {}".format(suffix))
    return {"sha256": sha256(path), "units": units, "warnings": warnings,
            "semantic_review": "not-performed"}


def extract_tree(source, output):
    output = outside_output(output, source)
    if inside(source, output):
        raise ValueError("Source and extraction output trees must not overlap")
    files, skipped = collect(source)
    output.mkdir(parents=True)
    results = []
    for path in files:
        rel = path.relative_to(source)
        item = {"source": rel.as_posix()}
        if path.suffix.lower() not in (".pdf", ".pptx", ".md", ".txt"):
            item.update(status="unsupported", reason="Use a suitable reader or visual inspection")
        else:
            try:
                data = extract_file(path)
                destination = output / (rel.as_posix() + ".json")
                write_new_json(destination, data)
                item.update(status="extracted", output=destination.relative_to(output).as_posix(),
                            empty_units=sum(u["empty_text"] for u in data["units"]))
            except Exception as exc:
                item.update(status="error", reason=str(exc))
        results.append(item)
    report = {"schema_version": 1, "results": results, "skipped": skipped,
              "ok": bool(results) and all(r["status"] == "extracted" for r in results)}
    write_new_json(output / "manifest.json", report)
    return report


def visible_markdown(text):
    """Exclude frontmatter, fenced and inline code from structural link checks."""
    text = re.sub(r"\A---\s*\n.*?\n---\s*\n", "", text, flags=re.S)
    result, fence = [], None
    for line in text.splitlines():
        match = re.match(r"^\s{0,3}(`{3,}|~{3,})", line)
        if match:
            marker = match.group(1)
            if fence is None:
                fence = marker
            elif marker[0] == fence[0] and len(marker) >= len(fence):
                fence = None
            continue
        if fence is None:
            result.append(re.sub(r"(`+).*?\1", "", line))
    return "\n".join(result)


def normal_heading(value):
    value = html.unescape(unquote(value))
    value = re.sub(r"[*_`~]", "", value).strip().rstrip("#").strip()
    return unicodedata.normalize("NFC", value).casefold()


def resolve_wiki(name, note, root, files):
    name = unquote(name)
    if not name:
        return [note]
    if name.startswith(("/", "\\")) or ".." in Path(name).parts:
        return []
    variants = [name] if Path(name).suffix else [name, name + ".md"]
    # Prefer an explicit vault path, then a note-relative path.
    for parent in (root, note.parent):
        candidates = [parent / v for v in variants if parent / v in files]
        if candidates:
            return candidates
    return [p for p in files if any(p.relative_to(root).as_posix() == v or
                                   p.relative_to(root).as_posix().endswith("/" + v)
                                   for v in variants)]


def check_links(root):
    listed, skipped = collect(root)
    files = set(listed)
    notes = {p: visible_markdown(p.read_text(encoding="utf-8-sig"))
             for p in listed if p.suffix.lower() == ".md"}
    errors, warnings, checked = [], [], 0
    for note, body in notes.items():
        rel = note.relative_to(root).as_posix()
        for match in re.finditer(r"!?\[\[([^\]\n]+)\]\]", body):
            raw = match.group(1).replace("\\|", "|").split("|", 1)[0]
            target, _, anchor = raw.partition("#")
            if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target):
                warnings.append({"note": rel, "link": raw, "reason": "External URI not checked"})
                continue
            checked += 1
            candidates = resolve_wiki(target, note, root, files)
            reason = None
            if len(candidates) != 1:
                reason = "Missing target" if not candidates else "Ambiguous target"
            elif anchor:
                destination = candidates[0]
                if destination in notes:
                    dest = notes[destination]
                    if anchor.startswith("^"):
                        if not re.search(r"(?:^|\s)" + re.escape(anchor) + r"\s*$", dest, flags=re.M):
                            reason = "Missing block ID"
                    else:
                        headings = [normal_heading(h) for h in re.findall(r"^#{1,6}\s+(.+)$", dest, re.M)]
                        if normal_heading(anchor) not in headings:
                            reason = "Missing heading (complex/nested heading syntax is unsupported)"
                elif destination.suffix.lower() == ".pdf" and re.fullmatch(r"page=[1-9]\d*", anchor):
                    try:
                        from pypdf import PdfReader
                    except ImportError:
                        warnings.append({"note": rel, "link": raw, "reason": "PDF page range unchecked: pypdf unavailable"})
                    else:
                        try:
                            if int(anchor[5:]) > len(PdfReader(str(destination)).pages):
                                reason = "PDF page out of range"
                        except Exception:
                            reason = "Cannot inspect PDF page range"
                else:
                    warnings.append({"note": rel, "link": raw, "reason": "Attachment fragment not checked"})
            if reason:
                errors.append({"note": rel, "link": raw, "reason": reason})
        if re.search(r"\]\(|\]\[", body):
            warnings.append({"note": rel, "reason": "Markdown-style links need a separate check"})
    return {"ok": not errors, "notes": len(notes), "wiki_links_checked": checked,
            "errors": errors, "warnings": warnings, "skipped": skipped,
            "limits": "Wiki links only; no math, source-fidelity or rendered-UI validation."}


def compare_snapshot(root, snapshot):
    before = json.loads(Path(snapshot).read_text(encoding="utf-8"))
    if before.get("schema_version") != 1 or not isinstance(before.get("files"), list):
        raise ValueError("Unsupported snapshot schema")
    old = {row["path"]: row["sha256"] for row in before["files"]}
    current = inventory(root)
    new = {row["path"]: row["sha256"] for row in current["files"]}
    changed = sorted(p for p in old.keys() & new.keys() if old[p] != new[p])
    missing = sorted(old.keys() - new.keys())
    return {"ok": not changed and not missing, "changed": changed, "missing": missing,
            "added": sorted(new.keys() - old.keys()), "skipped": current["skipped"]}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("inventory", "snapshot"):
        p = commands.add_parser(name)
        p.add_argument("--source", required=True)
        p.add_argument("--output", required=(name == "snapshot"))
    p = commands.add_parser("init")
    p.add_argument("--vault", required=True)
    p.add_argument("--course", required=True)
    p.add_argument("--title", required=True)
    p = commands.add_parser("extract")
    p.add_argument("--source", required=True)
    p.add_argument("--output", required=True)
    p = commands.add_parser("check")
    p.add_argument("--vault", required=True)
    p = commands.add_parser("compare")
    p.add_argument("--source", required=True)
    p.add_argument("--snapshot", required=True)
    args = parser.parse_args(argv)
    try:
        if args.command in ("inventory", "snapshot"):
            source = root_dir(args.source)
            output = outside_output(args.output, source) if args.output else None
            result = inventory(source)
            if output:
                write_new_json(output, result)
        elif args.command == "init":
            result = initialize(args.vault, args.course, args.title)
        elif args.command == "extract":
            result = extract_tree(root_dir(args.source), args.output)
        elif args.command == "check":
            result = check_links(root_dir(args.vault))
        else:
            result = compare_snapshot(root_dir(args.source), args.snapshot)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result.get("ok", True) else 1
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
